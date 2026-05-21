"""For chapters with Q# duplicates, surface the duplicated rows side by side."""
from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

EXCEL_DIR = Path(__file__).resolve().parents[2] / "data" / "Intermediate Financial Accounting test bank" / "excel"


def find_one(chapter: str) -> Path:
    matches = [p for p in EXCEL_DIR.iterdir()
               if p.suffix == ".xlsx" and not p.name.startswith("~$")
               and f"Chapter{chapter}" in p.name]
    if len(matches) != 1:
        raise RuntimeError(f"ch{chapter}: {matches}")
    return matches[0]


def dump_dupes(chapter: str) -> None:
    path = find_one(chapter)
    print("=" * 70)
    print(f"CH {chapter}: {path.name}")
    print("=" * 70)

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = list(rows[0])
    data = rows[1:]

    by_qnum: dict[str, list[tuple[int, tuple]]] = defaultdict(list)
    for i, raw in enumerate(data, start=2):
        qnum = raw[0]
        if qnum is None:
            continue
        by_qnum[str(qnum).strip()].append((i, raw))

    dupes = {q: rows_ for q, rows_ in by_qnum.items() if len(rows_) > 1}
    if not dupes:
        print("  no duplicates")
        return
    for qnum, rows_ in sorted(dupes.items(), key=lambda x: int(x[0]) if x[0].isdigit() else -1):
        print(f"\n  ---- Q# {qnum!r} appears {len(rows_)}x ----")
        for excel_row, raw in rows_:
            print(f"    [row {excel_row}]")
            for col_name, val in zip(header, raw):
                s = "" if val is None else str(val)
                if len(s) > 110:
                    s = s[:110] + "..."
                if s:
                    print(f"      {col_name!r:25s} {s!r}")


if __name__ == "__main__":
    for ch in ("05", "09"):
        dump_dupes(ch)
        print()
