"""Across all 21 chapters, count rows that look like multi-part 'fragments':
short Question text + missing core metadata (Bloom/Difficulty/LO)."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

EXCEL_DIR = Path(r"D:\Github\accounting\data\Intermediate Financial Accounting test bank\excel")


def find_one(chapter: str) -> Path:
    matches = [p for p in EXCEL_DIR.iterdir()
               if p.suffix == ".xlsx" and not p.name.startswith("~$")
               and f"Chapter{chapter}" in p.name]
    return matches[0]


def is_fragment(row: dict) -> bool:
    """A 'fragment' = a row that looks like a sub-part of an exploded multi-row
    problem, not a self-contained question."""
    q = (row.get("Question") or "").strip()
    bloom = (row.get("Bloom's") or "").strip()
    diff = (row.get("Difficulty") or "").strip()
    lo = (row.get("Learning Objective") or "").strip()
    # Skip header artifacts
    if str(row.get("Question_Number") or "").strip() == "0":
        return False
    if not q:
        return False
    # Heuristic: prompts < 40 chars AND missing at least two of {bloom, diff, lo}
    metadata_missing = sum(1 for v in (bloom, diff, lo) if not v)
    return len(q) < 40 and metadata_missing >= 2


def main() -> None:
    print(f"{'Ch':>3} | {'rows':>5} | {'fragments':>9} | {'fragment %':>10} | sample fragments")
    print("-" * 100)
    grand_rows = 0
    grand_frags = 0
    for i in range(1, 22):
        ch = f"{i:02d}"
        wb = openpyxl.load_workbook(find_one(ch), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        header = list(rows[0])
        data = rows[1:]
        n_rows = len(data)
        frags: list[tuple[int, str]] = []
        for excel_row, raw in enumerate(data, start=2):
            d = {col: (raw[i] if i < len(raw) else None) for i, col in enumerate(header)}
            if is_fragment(d):
                frags.append((excel_row, str(d.get("Question") or "").strip()))
        grand_rows += n_rows
        grand_frags += len(frags)
        sample = "; ".join(f"r{r}:{q[:30]!r}" for r, q in frags[:3])
        pct = 100 * len(frags) / n_rows if n_rows else 0
        print(f"{ch:>3} | {n_rows:>5} | {len(frags):>9} | {pct:>9.1f}% | {sample}")
    print("-" * 100)
    pct = 100 * grand_frags / grand_rows if grand_rows else 0
    print(f"TOT | {grand_rows:>5} | {grand_frags:>9} | {pct:>9.1f}% |")


if __name__ == "__main__":
    main()
