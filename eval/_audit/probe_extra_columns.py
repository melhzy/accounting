"""For chapters with extra columns, sample what's actually in them."""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

EXCEL_DIR = Path(__file__).resolve().parents[2] / "data" / "Intermediate Financial Accounting test bank" / "excel"

# (chapter, extra_col_name, expected_col_index, "compare-with-col-name", compare-col-index)
PROBES = [
    # Explanation (col 5) — present in ch06+. Compare with Answer (col 4).
    ("06", "Explanation", 5, "Answer", 4),
    ("10", "Explanation", 5, "Answer", 4),
    ("15", "Explanation", 5, "Answer", 4),
    # Question Type (col 12, with space) — ch07+. Compare with Question_Type (col 1).
    ("07", "Question Type", 12, "Question_Type", 1),
    ("18", "Question Type", 12, "Question_Type", 1),
    # Blooms (col 13) — ch08-09. Compare with Bloom's (col 7).
    ("08", "Blooms", 13, "Bloom's", 7),
    ("09", "Blooms", 13, "Bloom's", 7),
    # Option (col 13) — ch10-17. Compare with Options (col 3).
    ("10", "Option", 13, "Options", 3),
    ("15", "Option", 13, "Options", 3),
]


def load_columns(chapter: str, cols: list[int]) -> list[tuple[int, list]]:
    matches = [p for p in EXCEL_DIR.iterdir()
               if p.suffix == ".xlsx"
               and not p.name.startswith("~$")
               and f"Chapter{chapter}" in p.name]
    if len(matches) != 1:
        raise RuntimeError(f"ch {chapter}: {matches}")
    wb = openpyxl.load_workbook(matches[0], read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    data = rows[1:]
    return [(i + 2, [r[c] if c < len(r) else None for c in cols]) for i, r in enumerate(data)]


def main() -> None:
    for ch, extra_name, extra_idx, compare_name, compare_idx in PROBES:
        print("=" * 70)
        print(f"CH {ch}: extra column {extra_name!r} (idx {extra_idx}) "
              f"vs {compare_name!r} (idx {compare_idx})")
        print("=" * 70)

        cells = load_columns(ch, [extra_idx, compare_idx])

        n_total = len(cells)
        n_extra_nonempty = sum(1 for _, (e, _) in cells if e not in (None, "", " "))
        n_compare_nonempty = sum(1 for _, (_, c) in cells if c not in (None, "", " "))
        n_same = sum(1 for _, (e, c) in cells
                     if e is not None and c is not None
                     and str(e).strip() == str(c).strip()
                     and str(e).strip() != "")
        n_diff = sum(1 for _, (e, c) in cells
                     if e is not None and c is not None
                     and str(e).strip() != str(c).strip()
                     and str(e).strip() != "" and str(c).strip() != "")
        n_extra_only = sum(1 for _, (e, c) in cells
                           if e is not None and str(e).strip() != ""
                           and (c is None or str(c).strip() == ""))

        print(f"  total rows:                       {n_total}")
        print(f"  extra col nonempty:               {n_extra_nonempty}  ({100*n_extra_nonempty/n_total:.1f}%)")
        print(f"  compare col nonempty:             {n_compare_nonempty}")
        print(f"  same value (both nonempty):       {n_same}")
        print(f"  different value (both nonempty):  {n_diff}")
        print(f"  extra-only (compare empty):       {n_extra_only}")

        # Sample first 3 non-empty extras
        samples = [(r, e, c) for r, (e, c) in cells if e not in (None, "", " ")][:3]
        print("\n  sample rows where extra is nonempty:")
        for r, e, c in samples:
            e_s = str(e)[:120]
            c_s = str(c)[:120] if c is not None else "<None>"
            print(f"    [row {r}] {extra_name}={e_s!r}")
            print(f"              {compare_name}={c_s!r}")
        print()


if __name__ == "__main__":
    main()
