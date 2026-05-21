"""Profile one chapter: Excel shape + Word shape + side-by-side sample.

Usage: python profile_one.py 05
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
from docx import Document

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

REPO_ROOT = Path(__file__).resolve().parents[2]
ROOT = REPO_ROOT / "data" / "Intermediate Financial Accounting test bank"
EXCEL_DIR = ROOT / "excel"
WORD_DIR = ROOT / "word"


def find_chapter_file(directory: Path, chapter: str, ext: str) -> Path:
    matches = [p for p in directory.iterdir()
               if p.suffix == ext
               and not p.name.startswith("~$")
               and f"Chapter{chapter}" in p.name]
    if not matches:
        raise FileNotFoundError(f"no {ext} for chapter {chapter}")
    if len(matches) > 1:
        raise RuntimeError(f"multiple {ext} for chapter {chapter}: {matches}")
    return matches[0]


def profile_excel(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    data = rows[1:]
    out = {
        "path": str(path),
        "sheet": ws.title,
        "n_columns": len(header),
        "header": list(header),
        "n_rows_total": len(rows),
        "n_data_rows": len(data),
        "first_data_row": data[0] if data else None,
        "last_data_row": data[-1] if data else None,
    }
    wb.close()
    return out


def profile_word(path: Path) -> dict:
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs]
    nonempty = [p for p in paragraphs if p.strip()]
    q_marker = re.compile(r"^\s*(\d+)\)\s")
    question_starts = [p for p in nonempty if q_marker.match(p)]
    answer_lines = [p for p in nonempty if re.match(r"^\s*Answer:\s", p)]
    bloom_lines = [p for p in nonempty if re.match(r"^\s*Bloom['’]?s\s*:", p)]
    n_tables = len(doc.tables)
    return {
        "path": str(path),
        "n_paragraphs_total": len(paragraphs),
        "n_paragraphs_nonempty": len(nonempty),
        "n_question_starts": len(question_starts),
        "first_question": question_starts[0] if question_starts else None,
        "last_question": question_starts[-1] if question_starts else None,
        "n_answer_lines": len(answer_lines),
        "n_bloom_lines": len(bloom_lines),
        "n_tables": n_tables,
        "sample_first_15": nonempty[:15],
    }


def main(chapter: str) -> None:
    xlsx = find_chapter_file(EXCEL_DIR, chapter, ".xlsx")
    docx = find_chapter_file(WORD_DIR, chapter, ".docx")

    ex = profile_excel(xlsx)
    wd = profile_word(docx)

    print("=" * 70)
    print(f"CHAPTER {chapter}")
    print("=" * 70)

    print("\n--- EXCEL ---")
    print(f"file:        {ex['path']}")
    print(f"sheet:       {ex['sheet']}")
    print(f"n_columns:   {ex['n_columns']}")
    print(f"n_data_rows: {ex['n_data_rows']}")
    print(f"\nheader: {ex['header']}")
    print(f"\nfirst data row (truncated cells to 100 chars):")
    if ex["first_data_row"]:
        for col, val in zip(ex["header"], ex["first_data_row"]):
            s = "" if val is None else str(val)
            if len(s) > 100:
                s = s[:100] + "..."
            print(f"  {col!r}: {s!r}")
    print(f"\nlast data row (truncated cells to 100 chars):")
    if ex["last_data_row"]:
        for col, val in zip(ex["header"], ex["last_data_row"]):
            s = "" if val is None else str(val)
            if len(s) > 100:
                s = s[:100] + "..."
            print(f"  {col!r}: {s!r}")

    print("\n--- WORD ---")
    print(f"file:                  {wd['path']}")
    print(f"n_paragraphs_total:    {wd['n_paragraphs_total']}")
    print(f"n_paragraphs_nonempty: {wd['n_paragraphs_nonempty']}")
    print(f"n_question_starts:     {wd['n_question_starts']}")
    print(f"first question line:   {wd['first_question']!r}")
    print(f"last question line:    {wd['last_question']!r}")
    print(f"n_answer_lines:        {wd['n_answer_lines']}")
    print(f"n_bloom_lines:         {wd['n_bloom_lines']}")
    print(f"n_tables:              {wd['n_tables']}")
    print("\nsample (first 15 nonempty paragraphs):")
    for i, p in enumerate(wd["sample_first_15"]):
        s = p if len(p) <= 120 else p[:120] + "..."
        print(f"  [{i:02d}] {s!r}")

    print("\n--- CROSS-CHECK ---")
    print(f"excel data_rows:    {ex['n_data_rows']}")
    print(f"word q_starts:      {wd['n_question_starts']}")
    delta = ex["n_data_rows"] - wd["n_question_starts"]
    print(f"delta (excel-word): {delta:+d}")


if __name__ == "__main__":
    chapter = sys.argv[1] if len(sys.argv) > 1 else "05"
    main(chapter)
