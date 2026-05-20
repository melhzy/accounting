"""Extract every table from the Spiceland 9e Word .docx test bank and
stitch each one to the question it belongs to.

Walks the .docx body in document order so paragraphs and tables remain
interleaved, then attaches each table to the most-recent question marker
("N) ...") seen above it.

Outputs:
  eval/spiceland9e_tables.jsonl                    one record per table
  eval/_audit/tables_extraction_report.md          human-readable summary
  eval/_audit/tables_extraction_report.json        machine summary
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.document import Document as _Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

REPO_ROOT = Path(__file__).resolve().parents[1]
WORD_DIR = REPO_ROOT / "data" / "Intermediate Financial Accounting test bank" / "word"
OUT_JSONL = REPO_ROOT / "eval" / "spiceland9e_tables.jsonl"
OUT_DIR_AUDIT = REPO_ROOT / "eval" / "_audit"
OUT_MD = OUT_DIR_AUDIT / "tables_extraction_report.md"
OUT_JSON = OUT_DIR_AUDIT / "tables_extraction_report.json"
OUT_JSONL.parent.mkdir(parents=True, exist_ok=True)
OUT_DIR_AUDIT.mkdir(parents=True, exist_ok=True)

Q_MARKER = re.compile(r"^\s*(\d+)\)(\s|$)")


def iter_block_items(parent):
    """Yield Paragraph / Table objects in document order.

    The default doc.paragraphs / doc.tables lose interleaving."""
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise TypeError(f"unexpected parent type: {type(parent)}")
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def row_cells(row) -> list[str]:
    """Return the row's cell texts, deduplicating consecutive cells that
    share the same underlying <w:tc> (i.e. horizontally merged cells)."""
    out: list[str] = []
    last_tc_id: int | None = None
    for cell in row.cells:
        tc_id = id(cell._tc)
        if tc_id == last_tc_id:
            # Horizontally merged with previous cell; emit blank to keep grid
            out.append("")
            continue
        last_tc_id = tc_id
        text = cell.text.strip()
        out.append(text)
    return out


def render_markdown(table: Table) -> str:
    rows = [row_cells(r) for r in table.rows]
    if not rows:
        return ""
    # Pad rows to the same column count
    n_cols = max(len(r) for r in rows)
    rows = [r + [""] * (n_cols - len(r)) for r in rows]
    # Escape pipes and collapse internal newlines to spaces for markdown safety
    def esc(s: str) -> str:
        return s.replace("\n", " ").replace("|", r"\|").strip()
    header_row = "| " + " | ".join(esc(c) for c in rows[0]) + " |"
    sep = "| " + " | ".join(["---"] * n_cols) + " |"
    body = ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows[1:]]
    return "\n".join([header_row, sep] + body)


def find_one(chapter: str) -> Path:
    matches = [p for p in WORD_DIR.iterdir()
               if p.suffix == ".docx" and not p.name.startswith("~$")
               and f"Chapter{chapter}" in p.name]
    if len(matches) != 1:
        raise RuntimeError(f"ch {chapter}: found {matches}")
    return matches[0]


def extract_chapter(chapter: str) -> tuple[list[dict], dict]:
    path = find_one(chapter)
    doc = Document(str(path))

    chapter_int = int(chapter)
    records: list[dict] = []
    current_q: int | None = None
    table_idx_in_chapter = 0
    table_idx_under_question = 0
    last_paragraph_text = ""
    pending_after_table: dict | None = None  # to fill context_after on next paragraph
    q_marker_seen: set[int] = set()
    skipped_inq_enumerations = 0
    orphan_tables = 0

    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            # Resolve any context_after we owe
            if pending_after_table is not None and text:
                pending_after_table["context_after"] = text[:300]
                pending_after_table = None
            if not text:
                continue
            m = Q_MARKER.match(text)
            if m:
                n = int(m.group(1))
                # Only accept as a new question if it advances the counter
                # OR if this is the very first question marker we've seen.
                if current_q is None:
                    current_q = n
                    table_idx_under_question = 0
                    q_marker_seen.add(n)
                elif n > current_q:
                    current_q = n
                    table_idx_under_question = 0
                    q_marker_seen.add(n)
                else:
                    # In-question enumeration (1) ... 2) ... 3) ...) — ignore
                    skipped_inq_enumerations += 1
            last_paragraph_text = text
        elif isinstance(block, Table):
            n_rows = len(block.rows)
            cells_2d = [row_cells(r) for r in block.rows]
            n_cols = max((len(r) for r in cells_2d), default=0)
            # Skip degenerate tables (single empty cell)
            nonempty_cells = sum(1 for r in cells_2d for c in r if c.strip())
            if nonempty_cells == 0:
                table_idx_in_chapter += 1
                continue
            md = render_markdown(block)
            if current_q is None:
                orphan_tables += 1
                qnum_val: int | None = None
                parent_id: str | None = None
                q_id_suffix = "orphan"
            else:
                qnum_val = current_q
                parent_id = f"ch{chapter_int:02d}_q{current_q:04d}"
                q_id_suffix = f"q{current_q:04d}"
            table_idx_under_question += 1
            rec = {
                "table_id": f"ch{chapter_int:02d}_{q_id_suffix}_t{table_idx_under_question:02d}",
                "chapter": chapter_int,
                "question_number": qnum_val,
                "parent_question_id": parent_id,
                "table_idx_in_chapter": table_idx_in_chapter,
                "table_idx_under_question": table_idx_under_question,
                "n_rows": n_rows,
                "n_cols": n_cols,
                "n_cells_nonempty": nonempty_cells,
                "cells": cells_2d,
                "markdown": md,
                "context_before": last_paragraph_text[:300],
                "context_after": None,  # filled by the next paragraph
                "source_docx": path.name,
                "source_table_idx_in_docx": table_idx_in_chapter,
            }
            records.append(rec)
            pending_after_table = rec
            table_idx_in_chapter += 1

    stats = {
        "chapter": chapter_int,
        "n_tables_extracted": len(records),
        "orphan_tables": orphan_tables,
        "skipped_inq_enumerations": skipped_inq_enumerations,
        "n_unique_questions_with_tables": len({r["question_number"] for r in records if r["question_number"] is not None}),
        "qmax_seen": max(q_marker_seen, default=0),
    }
    return records, stats


def main() -> None:
    all_records: list[dict] = []
    per_chapter_stats: list[dict] = []
    tables_per_question_global: Counter = Counter()

    for i in range(1, 22):
        ch = f"{i:02d}"
        print(f"[ch{ch}] extracting...", flush=True)
        records, stats = extract_chapter(ch)
        all_records.extend(records)
        per_chapter_stats.append(stats)
        for r in records:
            if r["question_number"] is not None:
                tables_per_question_global[(r["chapter"], r["question_number"])] += 1

    # Write JSONL (sort for determinism)
    all_records.sort(key=lambda r: (r["chapter"],
                                    r["question_number"] if r["question_number"] is not None else -1,
                                    r["table_idx_in_chapter"]))
    with OUT_JSONL.open("w", encoding="utf-8") as fh:
        for r in all_records:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    print(f"wrote {OUT_JSONL}  ({len(all_records)} tables)")

    # Aggregate
    total_tables = len(all_records)
    total_orphans = sum(s["orphan_tables"] for s in per_chapter_stats)
    total_questions_with_tables = sum(s["n_unique_questions_with_tables"] for s in per_chapter_stats)
    # Top 10 questions by number of attached tables
    top_q = tables_per_question_global.most_common(10)

    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "total_tables_extracted": total_tables,
        "total_orphan_tables": total_orphans,
        "total_questions_with_at_least_one_table": total_questions_with_tables,
        "per_chapter": per_chapter_stats,
        "top10_questions_by_table_count": [
            {"chapter": c, "question_number": q, "n_tables": n}
            for (c, q), n in top_q
        ],
    }
    OUT_JSON.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {OUT_JSON}")

    # Human-readable summary
    md = [
        "# Word-table extraction report",
        "",
        f"_Generated: {summary['generated_at']}_",
        "",
        f"- Total tables extracted: **{total_tables:,}**",
        f"- Orphan tables (no preceding question marker): **{total_orphans}**",
        f"- Distinct questions with ≥1 table: **{total_questions_with_tables:,}**",
        "",
        "## Per chapter",
        "",
        "| Ch | tables extracted | orphans | in-question enum skipped | qmax | distinct q with tables |",
        "|----|------------------|---------|--------------------------|------|------------------------|",
    ]
    for s in per_chapter_stats:
        md.append(
            f"| {s['chapter']:02d} | {s['n_tables_extracted']} | {s['orphan_tables']} | "
            f"{s['skipped_inq_enumerations']} | {s['qmax_seen']} | {s['n_unique_questions_with_tables']} |"
        )
    md += [
        "",
        "## Top 10 questions by attached-table count",
        "",
        "| chapter | Q# | n_tables |",
        "|--------:|---:|---------:|",
    ]
    for (c, q), n in top_q:
        md.append(f"| {c:02d} | {q} | {n} |")
    OUT_MD.write_text("\n".join(md), encoding="utf-8")
    print(f"wrote {OUT_MD}")
    print()
    print("=== SUMMARY ===")
    print(json.dumps({k: v for k, v in summary.items() if k != "per_chapter"}, indent=2))


if __name__ == "__main__":
    main()
